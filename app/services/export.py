from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Mapping, Any
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


class ExportService:
    def build_xlsx(
        self,
        name: str,
        headers: list[str],
        rows: Iterable[Mapping[str, Any]],
        *,
        filters: Mapping[str, Any] | None = None,
        totals: Mapping[str, Any] | None = None,
        timezone_name: str | None = None,
    ) -> Path:
        export_dir = Path(__file__).resolve().parents[1] / "webapp" / "exports"
        export_dir.mkdir(exist_ok=True)
        stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
        safe_name = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in name).strip("_") or "export"
        path = export_dir / f"{safe_name}-{stamp}.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Данные"
        ws.append(headers)
        materialized = list(rows)
        for row in materialized:
            ws.append([row.get(h) for h in headers])

        if filters or totals or timezone_name:
            meta = wb.create_sheet("Параметры")
            meta.append(["Параметр", "Значение"])
            if timezone_name:
                meta.append(["Timezone", timezone_name])
            for k, v in (filters or {}).items():
                meta.append([k, v])
            for k, v in (totals or {}).items():
                meta.append([f"Итого: {k}", v])

        for col in range(1, ws.max_column + 1):
            values = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 200) + 1)]
            ws.column_dimensions[get_column_letter(col)].width = min(max(max(map(len, values), default=10) + 2, 10), 45)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(path)
        return path


export_service = ExportService()
