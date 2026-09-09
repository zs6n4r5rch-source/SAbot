from pathlib import Path

from openpyxl import load_workbook

from app.services.export import ExportService


def test_build_xlsx_creates_unique_filename_and_metadata(tmp_path: Path):
    service = ExportService(tmp_path)

    first = service.build_xlsx(
        "statistics/report",
        ["name", "amount"],
        [{"name": "A", "amount": 10}],
        filters={"days": 30},
        totals={"rows": 1},
        timezone_name="Europe/Moscow",
    )
    second = service.build_xlsx(
        "statistics/report",
        ["name"],
        [{"name": "B"}],
        timezone_name="Europe/Moscow",
    )

    assert first != second
    assert first.name.startswith("statistics_report-")
    assert first.suffix == ".xlsx"
    assert first.exists() and second.exists()

    workbook = load_workbook(first, read_only=True)
    assert workbook.sheetnames == ["Данные", "Параметры"]
    metadata = list(workbook["Параметры"].iter_rows(values_only=True))
    assert ("Timezone", "Europe/Moscow") in metadata
    assert ("days", 30) in metadata
    assert ("Итого: rows", 1) in metadata
